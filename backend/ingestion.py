# ============================================
# DOCUMENT INGESTION & PROCESSING PIPELINE
# ============================================
# Responsável por:
# 1. Upload seguro de arquivos
# 2. Extração de texto (PDF, DOCX, XLSX, TXT)
# 3. Fragmentação em chunks com overlap
# 4. Geração de embeddings vetoriais
# 5. Persistência em banco de dados

import os
import hashlib
import logging
from pathlib import Path
from typing import Tuple, List, Optional
from datetime import datetime
import pickle

import numpy as np
from sqlalchemy.orm import Session

# Document parsing
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook

# Embeddings
from openai import OpenAI

# Database
from models import Documento, Chunk, Embedding, Cliente
from database import get_db_session

logger = logging.getLogger(__name__)

# ============================================
# CONSTANTS & CONFIGURATION
# ============================================

CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", 1000))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", 200))
ALLOWED_FILE_TYPES = os.getenv("ALLOWED_FILE_TYPES", "pdf,docx,xlsx,txt,doc,xls").split(",")
MAX_FILE_SIZE = int(os.getenv("MAX_FILE_SIZE_MB", 50)) * 1024 * 1024
UPLOAD_DIR = os.getenv("UPLOAD_DIR", "./uploads")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
EMBEDDING_MODEL = os.getenv("OPENAI_EMBEDDING_MODEL", "text-embedding-ada-002")

# Criar diretório de uploads se não existir
Path(UPLOAD_DIR).mkdir(parents=True, exist_ok=True)

# Cliente OpenAI
client = OpenAI(api_key=OPENAI_API_KEY)

# ============================================
# FILE VALIDATION
# ============================================

def validate_file(
    filename: str,
    file_size: int
) -> Tuple[bool, str]:
    """
    Valida nome e tamanho do arquivo antes do upload.
    
    Args:
        filename: Nome do arquivo
        file_size: Tamanho em bytes
        
    Returns:
        (is_valid, error_message)
    """
    # Validar extensão
    file_ext = filename.lower().split(".")[-1]
    if file_ext not in ALLOWED_FILE_TYPES:
        return False, f"Tipo de arquivo não permitido: .{file_ext}"
    
    # Validar tamanho
    if file_size > MAX_FILE_SIZE:
        max_mb = MAX_FILE_SIZE / (1024 * 1024)
        return False, f"Arquivo excede o limite de {max_mb}MB"
    
    return True, ""


def calculate_file_hash(file_path: str) -> str:
    """
    Calcula SHA-256 do arquivo para evitar duplicatas.
    """
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            sha256.update(chunk)
    return sha256.hexdigest()


# ============================================
# TEXT EXTRACTION BY FILE TYPE
# ============================================

def extract_text_from_pdf(file_path: str) -> Tuple[str, int]:
    """
    Extrai texto de arquivo PDF.
    
    Returns:
        (texto_completo, numero_paginas)
    """
    try:
        logger.info(f"Extracting text from PDF: {file_path}")
        reader = PdfReader(file_path)
        text_parts = []
        
        for page_num, page in enumerate(reader.pages):
            try:
                text = page.extract_text()
                if text:
                    # Adicionar marcador de página para tracking
                    text_parts.append(f"[PAGE {page_num + 1}]\n{text}")
            except Exception as e:
                logger.warning(f"Failed to extract text from page {page_num + 1}: {e}")
        
        full_text = "\n".join(text_parts)
        num_pages = len(reader.pages)
        
        logger.info(f"Extracted {len(full_text)} characters from {num_pages} pages")
        return full_text, num_pages
        
    except Exception as e:
        logger.error(f"Error extracting PDF text: {e}")
        raise


def extract_text_from_docx(file_path: str) -> Tuple[str, int]:
    """
    Extrai texto de arquivo DOCX (Word).
    
    Returns:
        (texto_completo, numero_paragrafos)
    """
    try:
        logger.info(f"Extracting text from DOCX: {file_path}")
        doc = DocxDocument(file_path)
        text_parts = []
        
        for para_num, para in enumerate(doc.paragraphs):
            if para.text.strip():
                text_parts.append(para.text)
        
        # Tabelas também
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_text.append(" | ".join(row_data))
            text_parts.append("\n".join(table_text))
        
        full_text = "\n".join(text_parts)
        num_paragraphs = len(doc.paragraphs)
        
        logger.info(f"Extracted {len(full_text)} characters from {num_paragraphs} paragraphs")
        return full_text, num_paragraphs
        
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        raise


def extract_text_from_xlsx(file_path: str) -> Tuple[str, int]:
    """
    Extrai texto de arquivo XLSX (Excel).
    Inclui nome da aba, cabeçalhos e dados.
    
    Returns:
        (texto_completo, numero_linhas)
    """
    try:
        logger.info(f"Extracting text from XLSX: {file_path}")
        workbook = load_workbook(file_path)
        text_parts = []
        total_rows = 0
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text_parts.append(f"[SHEET: {sheet_name}]")
            
            # Extrair dados em formato tabular
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    # Filtrar e formatar células
                    row_data = [str(cell) if cell is not None else "" for cell in row]
                    text_parts.append(" | ".join(row_data))
                    total_rows += 1
            
            text_parts.append("")  # Separador entre abas
        
        full_text = "\n".join(text_parts)
        
        logger.info(f"Extracted {len(full_text)} characters from {total_rows} rows")
        return full_text, total_rows
        
    except Exception as e:
        logger.error(f"Error extracting XLSX text: {e}")
        raise


def extract_text_from_txt(file_path: str) -> Tuple[str, int]:
    """
    Lê arquivo TXT simples.
    
    Returns:
        (texto_completo, numero_linhas)
    """
    try:
        logger.info(f"Extracting text from TXT: {file_path}")
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
        
        num_lines = len(text.split("\n"))
        
        logger.info(f"Extracted {len(text)} characters from {num_lines} lines")
        return text, num_lines
        
    except Exception as e:
        logger.error(f"Error extracting TXT text: {e}")
        raise


def extract_text(file_path: str, file_type: str) -> Tuple[str, int]:
    """
    Dispatcher para extrair texto baseado no tipo de arquivo.
    """
    file_type_lower = file_type.lower()
    
    if file_type_lower == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type_lower in ["docx", "doc"]:
        return extract_text_from_docx(file_path)
    elif file_type_lower in ["xlsx", "xls"]:
        return extract_text_from_xlsx(file_path)
    elif file_type_lower == "txt":
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


# ============================================
# TEXT CHUNKING
# ============================================

def create_chunks(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """
    Fragmenta texto em chunks com overlap.
    Estratégia: dividir por caracteres com sobreposição para preservar contexto.
    
    Args:
        text: Texto completo
        chunk_size: Tamanho de cada chunk em caracteres
        overlap: Número de caracteres sobrepostos entre chunks
        
    Returns:
        Lista de chunks
    """
    if len(text) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    
    while start < len(text):
        # Calcular fim do chunk
        end = min(start + chunk_size, len(text))
        
        # Tentar quebrar em espaço/nova linha para não cortar palavras
        if end < len(text):
            # Procurar última quebra de linha antes do fim
            last_newline = text.rfind("\n", start, end)
            if last_newline > start:
                end = last_newline + 1
            else:
                # Procurar último espaço
                last_space = text.rfind(" ", start, end)
                if last_space > start:
                    end = last_space + 1
        
        chunk = text[start:end].strip()
        if chunk:  # Ignorar chunks vazios
            chunks.append(chunk)
        
        # Mover start para próximo chunk com overlap
        start = end - overlap
    
    logger.info(f"Created {len(chunks)} chunks from text")
    return chunks


# ============================================
# EMBEDDING GENERATION
# ============================================

def generate_embeddings_batch(texts: List[str], model: str = EMBEDDING_MODEL) -> List[np.ndarray]:
    """
    Gera embeddings para uma lista de textos usando OpenAI API.
    Agrupa em batches para economizar tempo/custo.
    
    Args:
        texts: Lista de textos para gerar embeddings
        model: Modelo de embedding (padrão: text-embedding-ada-002)
        
    Returns:
        Lista de arrays numpy com embeddings
    """
    if not texts:
        return []
    
    try:
        logger.info(f"Generating embeddings for {len(texts)} texts using {model}")
        
        # OpenAI recomenda não enviar mais de 2048 textos por requisição
        batch_size = 2048
        all_embeddings = []
        
        for i in range(0, len(texts), batch_size):
            batch = texts[i:i + batch_size]
            
            response = client.embeddings.create(
                input=batch,
                model=model
            )
            
            # Extrair embeddings da resposta
            for item in response.data:
                embedding = np.array(item.embedding, dtype=np.float32)
                all_embeddings.append(embedding)
            
            logger.info(f"Generated embeddings batch {i//batch_size + 1}")
        
        logger.info(f"Successfully generated {len(all_embeddings)} embeddings")
        return all_embeddings
        
    except Exception as e:
        logger.error(f"Error generating embeddings: {e}")
        raise


# ============================================
# DATABASE PERSISTENCE
# ============================================

def save_document_to_db(
    db: Session,
    cliente_id: int,
    file_path: str,
    filename: str,
    file_type: str,
    conteudo_texto: str,
    hash_arquivo: str,
    tamanho_bytes: int
) -> Documento:
    """
    Salva metadados do documento no banco de dados.
    """
    documento = Documento(
        cliente_id=cliente_id,
        nome_arquivo=filename,
        tipo_arquivo=file_type.upper(),
        caminho_arquivo=file_path,
        tamanho_bytes=tamanho_bytes,
        hash_arquivo=hash_arquivo,
        conteudo_texto=conteudo_texto,
        status="PROCESSANDO"
    )
    db.add(documento)
    db.commit()
    db.refresh(documento)
    
    logger.info(f"Documento {documento.id} salvo no banco de dados")
    return documento


def save_chunks_to_db(
    db: Session,
    documento_id: int,
    chunks: List[str],
    page_info: Optional[dict] = None
) -> List[Chunk]:
    """
    Salva chunks no banco de dados.
    """
    chunk_objects = []
    
    for numero, conteudo in enumerate(chunks):
        pagina = None
        if page_info and numero in page_info:
            pagina = page_info[numero]
        
        chunk = Chunk(
            documento_id=documento_id,
            numero_chunk=numero,
            conteudo=conteudo,
            tamanho_caracteres=len(conteudo),
            pagina=pagina
        )
        chunk_objects.append(chunk)
    
    db.bulk_save_objects(chunk_objects)
    db.commit()
    
    logger.info(f"Salvos {len(chunk_objects)} chunks para documento {documento_id}")
    return chunk_objects


def save_embeddings_to_db(
    db: Session,
    chunks: List[Chunk],
    embeddings: List[np.ndarray],
    modelo: str = EMBEDDING_MODEL
) -> int:
    """
    Salva embeddings (vetores) no banco de dados.
    Embeddings são serializados como pickle em BLOB.
    """
    embedding_objects = []
    
    for chunk, embedding in zip(chunks, embeddings):
        # Serializar embedding como pickle
        embedding_pickle = pickle.dumps(embedding)
        
        embedding_obj = Embedding(
            chunk_id=chunk.id,
            vetor=embedding_pickle,
            dimensao=len(embedding),
            modelo=modelo
        )
        embedding_objects.append(embedding_obj)
    
    db.bulk_save_objects(embedding_objects)
    db.commit()
    
    logger.info(f"Salvos {len(embedding_objects)} embeddings")
    return len(embedding_objects)


def mark_document_complete(db: Session, documento_id: int, total_chunks: int, total_embeddings: int):
    """
    Marca documento como processado com sucesso.
    """
    documento = db.query(Documento).filter(Documento.id == documento_id).first()
    if documento:
        documento.status = "SUCESSO"
        documento.data_processamento = datetime.utcnow()
        documento.total_chunks = total_chunks
        documento.total_embeddings = total_embeddings
        db.commit()
        logger.info(f"Documento {documento_id} marcado como SUCESSO")


def mark_document_error(db: Session, documento_id: int, error_message: str):
    """
    Marca documento como com erro.
    """
    documento = db.query(Documento).filter(Documento.id == documento_id).first()
    if documento:
        documento.status = "ERRO"
        documento.mensagem_erro = error_message
        db.commit()
        logger.error(f"Documento {documento_id} marcado como ERRO: {error_message}")


# ============================================
# MAIN INGESTION PIPELINE
# ============================================

def process_document(
    cliente_id: int,
    file_path: str,
    filename: str,
    file_type: str
) -> Tuple[bool, str, Optional[Documento]]:
    """
    Pipeline completo de processamento de documento:
    1. Extração de texto
    2. Criação de chunks
    3. Geração de embeddings
    4. Persistência em banco de dados
    
    Returns:
        (sucesso, mensagem, documento)
    """
    documento = None
    
    try:
        logger.info(f"Starting document processing: {filename}")
        
        with get_db_session() as db:
            # 1. Validar cliente existe
            cliente = db.query(Cliente).filter(Cliente.id == cliente_id).first()
            if not cliente:
                return False, f"Cliente {cliente_id} não encontrado", None
            
            # 2. Calcular hash e verificar duplicata
            hash_arquivo = calculate_file_hash(file_path)
            existing = db.query(Documento).filter(
                Documento.hash_arquivo == hash_arquivo
            ).first()
            if existing:
                logger.warning(f"Documento duplicado detectado: {hash_arquivo}")
                return False, "Este documento já foi processado", None
            
            # 3. Obter tamanho do arquivo
            tamanho_bytes = os.path.getsize(file_path)
            
            # 4. Extrair texto
            try:
                conteudo_texto, num_pages = extract_text(file_path, file_type)
                if len(conteudo_texto.strip()) < 100:
                    return False, "Documento contém texto insuficiente", None
            except Exception as e:
                return False, f"Erro ao extrair texto: {str(e)}", None
            
            # 5. Salvar documento no BD (status PROCESSANDO)
            documento = save_document_to_db(
                db, cliente_id, file_path, filename, file_type,
                conteudo_texto, hash_arquivo, tamanho_bytes
            )
            documento_id = documento.id
        
        # 6. Fragmentar texto
        chunks = create_chunks(conteudo_texto)
        
        # 7. Gerar embeddings
        embeddings = generate_embeddings_batch(chunks)
        
        if len(embeddings) != len(chunks):
            raise ValueError("Número de embeddings não corresponde a chunks")
        
        # 8. Salvar chunks e embeddings no BD
        with get_db_session() as db:
            chunk_objects = save_chunks_to_db(db, documento_id, chunks)
            embeddings_count = save_embeddings_to_db(db, chunk_objects, embeddings)
            mark_document_complete(db, documento_id, len(chunk_objects), embeddings_count)
        
        logger.info(f"Document processing completed successfully: {filename}")
        return True, "Documento processado com sucesso", documento
        
    except Exception as e:
        logger.error(f"Document processing failed: {e}")
        if documento:
            with get_db_session() as db:
                mark_document_error(db, documento.id, str(e))
        
        return False, f"Erro ao processar documento: {str(e)}", documento


def get_document_status(documento_id: int) -> dict:
    """
    Retorna status atual de um documento.
    """
    with get_db_session() as db:
        documento = db.query(Documento).filter(Documento.id == documento_id).first()
        if not documento:
            return {"erro": "Documento não encontrado"}
        
        return {
            "id": documento.id,
            "nome": documento.nome_arquivo,
            "status": documento.status,
            "total_chunks": documento.total_chunks,
            "total_embeddings": documento.total_embeddings,
            "tamanho_bytes": documento.tamanho_bytes,
            "data_upload": documento.data_upload.isoformat() if documento.data_upload else None,
            "data_processamento": documento.data_processamento.isoformat() if documento.data_processamento else None,
            "mensagem_erro": documento.mensagem_erro
        }
